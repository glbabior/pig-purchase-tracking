from openpyxl import load_workbook
from pathlib import Path

path = Path('MonthlyBudget.xlsx')
wb = load_workbook(path, data_only=True)
print('sheets', wb.sheetnames)
ws = wb.active
for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
    print(idx, row)
    if idx >= 40:
        break
